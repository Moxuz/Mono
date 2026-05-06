"""
Fix remaining stale test result values in TAS8-final (1).docx
"""
import sys
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r"C:\Users\ASUS\Documents\TEST\TESTMONO\Final-Final\TAS8-final (1).docx"
doc = Document(DOC_PATH)

def set_cell_text(cell, new_text):
    para = cell.paragraphs[0]
    if para.runs:
        run0 = para.runs[0]
        bold, italic, size, name = run0.bold, run0.italic, run0.font.size, run0.font.name
        for run in para.runs:
            run.text = ''
        run0.text = new_text
        run0.bold, run0.italic = bold, italic
        if size: run0.font.size = size
        if name: run0.font.name = name
    else:
        para.add_run(new_text)

def replace_para(para, replacements):
    """Apply list of (old, new) substitutions to paragraph, rebuilding first run."""
    text = para.text
    changed = False
    for old, new in replacements:
        if old in text:
            text = text.replace(old, new)
            changed = True
    if changed and para.runs:
        bold, italic, size, name = para.runs[0].bold, para.runs[0].italic, para.runs[0].font.size, para.runs[0].font.name
        for run in para.runs:
            run.text = ''
        para.runs[0].text = text
        para.runs[0].bold, para.runs[0].italic = bold, italic
        if size: para.runs[0].font.size = size
        if name: para.runs[0].font.name = name
    return changed

paras = doc.paragraphs
changes = []

# ── [920] Table caption: Jest 104 → 108 ───────────────────────────────────────
if replace_para(paras[920], [('(104 test cases)', '(108 test cases)')]):
    changes.append('[920] Table caption Jest: 104→108')

# ── [1083] Summary bullet: Jest 104 → 108 ────────────────────────────────────
if replace_para(paras[1083], [
    ('104 test cases, ผ่านทั้งหมด 104/104', '108 test cases, ผ่านทั้งหมด 108/108'),
    ('104 test cases, ผ่าน 104/104', '108 test cases, ผ่าน 108/108'),
]):
    changes.append('[1083] Summary Jest: 104→108')

# ── [1084] Summary bullet: Newman ────────────────────────────────────────────
if replace_para(paras[1084], [
    ('70 assertions, ผ่าน 69/70 (98.6%) — 1 assertion ไม่ผ่าน (logout 504 timeout, known issue)',
     '69 assertions, ผ่าน 69/69 (100.0%) — ผ่านทุก assertion'),
    ('70 assertions, ผ่าน 69/70 (98.6%)',
     '69 assertions, ผ่าน 69/69 (100.0%)'),
]):
    changes.append('[1084] Summary Newman: 70→69 assertions, 98.6%→100.0%')

# ── [1085] Summary bullet: Playwright ────────────────────────────────────────
if replace_para(paras[1085], [
    ('186 test cases, ผ่าน 182/186 (97.8%), 4 skip (OAuth state dependency)',
     '211 test cases, ผ่าน 206/211 (97.6%), 38 กลุ่ม, 3 skip (OAuth state), 2 flaky (timing)'),
    ('186 test cases, ผ่าน 182/186 (97.8%)',
     '211 test cases, ผ่าน 206/211 (97.6%)'),
]):
    changes.append('[1085] Summary Playwright: 186→211, 182→206, 97.8%→97.6%')

# ── [1087] Summary bullet: ZAP ───────────────────────────────────────────────
if replace_para(paras[1087], [
    ('55 PASS, 4 Medium (CSP), 4 Low, 9 Informational',
     '59 PASS, 5 Medium (CSP), 1 Low, 6 Informational'),
    ('0 FAIL, 55 PASS',
     '0 FAIL, 59 PASS'),
]):
    changes.append('[1087] Summary ZAP: 55→59 PASS, 4 Medium→5, 4 Low→1, 9 Info→6')

# ── [1099] Context paragraph: Jest 104 ───────────────────────────────────────
# Only if "104 test cases" appears clearly referring to Jest unit tests
p1099 = paras[1099]
if '104 test cases' in p1099.text or '104/104' in p1099.text:
    if replace_para(p1099, [
        ('104 test cases', '108 test cases'),
        ('104/104', '108/108'),
    ]):
        changes.append('[1099] Conclusion paragraph: 104→108')

# ── TABLE[14] — Newman detail table ──────────────────────────────────────────
# Row[2] Authentication: assertions 22→21, pass 21→21, fail 1*→0
t14 = doc.tables[14]
# Row[2]: Group, Requests, Assertions, Pass, Fail
t14_r2 = t14.rows[2]
cells_r2 = [c.text.strip() for c in t14_r2.cells]
print(f"Table[14] Row[2] before: {cells_r2}")
set_cell_text(t14_r2.cells[2], '21')    # assertions: 22→21 (removed the 504 one)
set_cell_text(t14_r2.cells[3], '21')    # pass: stays 21
set_cell_text(t14_r2.cells[4], '0')     # fail: 1*→0
changes.append('Table[14] Row[2] Auth: assertions 22→21, fail 1*→0')

# Row[8] TOTAL: assertions 70→69, fail 1*→0
t14_r8 = t14.rows[8]
cells_r8 = [c.text.strip() for c in t14_r8.cells]
print(f"Table[14] Row[8] before: {cells_r8}")
set_cell_text(t14_r8.cells[2], '69')    # assertions: 70→69
set_cell_text(t14_r8.cells[3], '69')    # pass: stays 69
set_cell_text(t14_r8.cells[4], '0')     # fail: 1*→0
changes.append('Table[14] Row[8] TOTAL: assertions 70→69, fail 1*→0')

# ── Save ─────────────────────────────────────────────────────────────────────
SAVE_PATH = DOC_PATH.replace('.docx', '_updated.docx')
doc.save(SAVE_PATH)
print(f"\nSaved to: {SAVE_PATH}")
print("\n=== Changes applied ===")
for c in changes:
    print(f"  OK: {c}")
print(f"\nTotal changes: {len(changes)}")
