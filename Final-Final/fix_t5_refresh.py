"""
Fix Table[5] Row[2] — refresh token expiry says "1 hr (OAuth)" but code uses 30d for OAuth too
"""
import sys, shutil
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')
DOC_PATH = r"C:\Users\ASUS\Documents\TEST\TESTMONO\Final-Final\TAS8-final (1).docx"
BAK_PATH = r"C:\Users\ASUS\Documents\TEST\TESTMONO\Final-Final\TAS8-final (1).docx.bak_t5fix"

shutil.copy2(DOC_PATH, BAK_PATH)
doc = Document(DOC_PATH)

t5 = doc.tables[5]
r2 = t5.rows[2]

print('Before:')
for ci, cell in enumerate(r2.cells):
    print(f'  Col[{ci}]: "{cell.text.strip()}"')

def fix_cell(cell, old, new, label=''):
    for para in cell.paragraphs:
        full = ''.join(r.text for r in para.runs)
        if old in full:
            if para.runs:
                para.runs[0].text = full.replace(old, new, 1)
                for r in para.runs[1:]:
                    r.text = ''
            else:
                para.add_run(full.replace(old, new, 1))
            print(f'  OK  {label}: "{old}" → "{new}"')
            return True
    print(f'  SKP {label}: "{old}" not found')
    return False

# Fix Col[2]: "30 days (session) / 1 hr (OAuth)" → "30 days (ทั้ง session และ OAuth)"
fix_cell(r2.cells[2],
    '30 days (session) / 1 hr (OAuth)',
    '30 days (ทั้ง session refresh token และ OAuth refresh token)',
    'T5R2C2-refresh-expiry')

# Fix Col[3] remark: "OAuth refresh token ตั้งใจให้มีอายุสั้นกว่า" → correct note
fix_cell(r2.cells[3],
    'OAuth refresh token ตั้งใจให้มีอายุสั้นกว่า',
    'OAuth refresh token มีอายุ 30 วัน เช่นเดียวกับ session refresh token',
    'T5R2C3-remark')

print('\nAfter:')
for ci, cell in enumerate(r2.cells):
    print(f'  Col[{ci}]: "{cell.text.strip()}"')

# Verify
actual = t5.rows[2].cells[2].text.strip()
assert '1 hr (OAuth)' not in actual, f'Old value still present: {actual}'
assert '30 days' in actual, f'30 days not found: {actual}'
print('\nVERIFY: PASS')

doc.save(DOC_PATH)
print(f'Saved: {DOC_PATH}')
