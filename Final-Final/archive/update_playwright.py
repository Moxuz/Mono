"""
update_playwright.py — adds Playwright E2E test results §4.2.6 to TAS6-final.docx
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt

doc = Document('TAS6-final.docx')

def insert_after(ref_para, text, bold=False, italic=False):
    new_p = OxmlElement('w:p')
    ref_para._p.addnext(new_p)
    paras = doc.paragraphs
    ref_idx = next(i for i, p in enumerate(paras) if p._p is ref_para._p)
    new_para = paras[ref_idx + 1]
    if text:
        run = new_para.add_run(text)
        run.bold = bold
        run.italic = italic
    return new_para

# Find anchor: end of Postman/Newman section
anchor = None
for para in doc.paragraphs:
    if 'Newman v6.2.2' in para.text and 'Known issue' in para.text:
        anchor = para
        print(f"Found Postman anchor: {para.text[:70]}")
        break

if not anchor:
    # Fallback: find by known issue text
    for para in doc.paragraphs:
        if 'logout timeout' in para.text and '504' in para.text and 'documented' in para.text:
            anchor = para
            print(f"Fallback anchor: {para.text[:70]}")
            break

if not anchor:
    print("ERROR: anchor not found"); exit(1)

print("Inserting §4.2.6 Playwright E2E Testing...")
cur = anchor

cur = insert_after(cur, '')

cur = insert_after(cur,
    '4.2.6 End-to-End Testing (Playwright)',
    bold=True)

cur = insert_after(cur,
    'ทำการทดสอบแบบ End-to-End (E2E) ด้วย Playwright v1.x บน Chromium browser '
    'ครอบคลุมทั้ง UI interactions และ API calls ผ่าน browser context '
    'ทดสอบ 8 กลุ่ม (test describe) รวม 25 test cases '
    'รวมถึง page navigation, security headers, login/register form flow, '
    'OAuth 2.0, session management และ dashboard access control')

cur = insert_after(cur, '')

cur = insert_after(cur,
    'ตารางที่ 4.5 ผลการทดสอบ E2E ด้วย Playwright (Chromium)',
    bold=True)

rows = [
    ('Test Group',                        'Tests', 'Pass', 'Fail', 'Duration'),
    ('01 — Page Navigation',              '6',     '6',    '0',    '~14s'),
    ('02 — Security Headers',             '4',     '4',    '0',    '~5s'),
    ('03 — Login Flow',                   '3',     '3',    '0',    '~15s'),
    ('04 — Register Flow',                '2',     '2',    '0',    '~14s'),
    ('05 — API Calls via Browser (fetch)','5',     '5',    '0',    '~4s'),
    ('06 — OAuth 2.0 via Browser',        '2',     '2',    '0',    '~3s'),
    ('07 — Session Management',           '2',     '2',    '0',    '~2s'),
    ('08 — Dashboard Access Control',     '1',     '1',    '0',    '~1s'),
    ('TOTAL',                             '25',    '25',   '0',    '21.8s'),
]
for row in rows:
    cur = insert_after(cur, ' | '.join(row))

cur = insert_after(cur, '')

cur = insert_after(cur,
    'ผลการทดสอบสำคัญ:')

findings = [
    'Page Navigation: ทุกหน้า (/, /login, /register, /forgot-password, /dashboard) โหลดสำเร็จ; 404 ไม่ทำให้ระบบ crash',
    'Security Headers: CSP, X-Content-Type-Options (nosniff), X-XSS-Protection ครบทุก response; ไม่มี X-Powered-By',
    'Health endpoint: ตอบสนองใน 413ms พร้อม status: OK และ uptime ถูกต้อง',
    'Login Form: ส่งข้อมูลถูกต้อง redirect สำเร็จ; password ผิดแสดง error; email ว่างไม่ผ่าน validation',
    'Register Form: ตรวจพบ duplicate email; มี fields ครบ (email, password, confirmPassword, consent)',
    'API via fetch: Login, Profile, Refresh-token, OIDC discovery ทำงานถูกต้องในบริบท browser',
    'OAuth: Client registration และ userinfo endpoint ทำงานได้จาก browser context',
    'Sessions: GET /api/sessions และ revoke-all-others ทำงานถูกต้อง',
    'Access Control: Dashboard endpoints ทั้ง 3 ถูกปฏิเสธ (401) พร้อมกัน (parallel fetch)',
]
for f in findings:
    cur = insert_after(cur, f'- {f}')

cur = insert_after(cur, '')
cur = insert_after(cur,
    'เครื่องมือที่ใช้: Playwright v1.x, Browser: Chromium (headless), '
    'จำนวน test cases: 25, ผล: 25/25 PASS (100%), '
    'เวลารวม: 21.8s (parallel execution ด้วย 8 workers)',
    italic=True)

doc.save('TAS6-final.docx')
print("Saved TAS6-final.docx with §4.2.6 Playwright results")
