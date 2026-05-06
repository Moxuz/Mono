import sys
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r"C:\Users\ASUS\Documents\TEST\TESTMONO\Final-Final\TAS8-final (1).docx"
doc = Document(DOC_PATH)

print(f"Total tables: {len(doc.tables)}")
print(f"Total paragraphs: {len(doc.paragraphs)}")

# Print ALL tables (all rows)
for ti, table in enumerate(doc.tables):
    print(f"\n{'='*70}")
    print(f"TABLE[{ti}] — {len(table.rows)} rows x {len(table.columns)} cols")
    print(f"{'='*70}")
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip()[:60] for c in row.cells]
        print(f"  Row[{ri}]: {cells}")

# Scan ALL paragraphs for any old test numbers still present
print(f"\n{'='*70}")
print("SCANNING ALL PARAGRAPHS for old values...")
print(f"{'='*70}")
old_values = ['104', '70 assertions', '69/70', '98.6', '186', '182', '55 PASS', '55 passive', '29 ม', '27 กลุ่ม', '141 test', '4 skip', '~1.5', '8 workers', '504 Gateway']
for i, p in enumerate(doc.paragraphs):
    txt = p.text
    if not txt.strip():
        continue
    for val in old_values:
        if val in txt:
            print(f"  [{i}] FOUND '{val}': {txt[:150]}")
            break
