"""
Full document audit v2 — checks everything after all fixes
Covers: all tables, technical facts, API routes, test results, placeholders, cross-refs
"""
import sys, re
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')
DOC_PATH = r"C:\Users\ASUS\Documents\TEST\TESTMONO\Final-Final\TAS8-final (1).docx"
doc = Document(DOC_PATH)
paras = doc.paragraphs

issues = []
ok_list = []
notes = []

def flag(label, detail): issues.append(f'  FAIL [{label}]: {detail}')
def good(label, detail=''): ok_list.append(f'  OK   [{label}]: {detail}')
def note(label, detail): notes.append(f'  NOTE [{label}]: {detail}')

# ─────────────────────────────────────────────────────────────────────────────
# 1. PLACEHOLDER / COMMENT SCAN
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 1. PLACEHOLDERS / COMMENTS ===')
placeholder_patterns = [r'TODO', r'FIXME', r'X\.X', r'\[FILL', r'\[TBD', r'TBD', r'ยังไม่']
# Note: // excluded because https:// is valid
found_any = False
for i, p in enumerate(paras):
    t = p.text
    if not t.strip(): continue
    for pat in placeholder_patterns:
        if re.search(pat, t, re.IGNORECASE):
            flag(f'para[{i}]', f'Placeholder "{pat}": {t[:120]}')
            found_any = True
            break
if not found_any:
    good('Placeholders', 'None found')

# ─────────────────────────────────────────────────────────────────────────────
# 2. STALE TECHNICAL VALUES IN PARAGRAPHS
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 2. STALE TECHNICAL VALUES IN PARAGRAPHS ===')
stale_patterns = [
    ('RS256 in body', r'RS256', 'HS256 is correct algorithm'),
    ('JWT_PRIVATE_KEY in body', r'JWT_PRIVATE_KEY', 'should be JWT_SECRET'),
    ('56 endpoints', r'56 endpoint', 'should be 44'),
    ('cost-12 in body', r'cost.?12|bcrypt.?12|12.?round', 'should be cost-10'),
    ('v1.x placeholder', r'v1\.x\b', 'should be specific version'),
    ('old refresh path', r'/api/auth/refresh(?!-token)', 'should be /api/auth/refresh-token'),
]
for label, pat, hint in stale_patterns:
    hits = []
    for i, p in enumerate(paras):
        t = p.text
        if not t.strip() or p.style.name.startswith('Heading'): continue
        if re.search(pat, t, re.IGNORECASE):
            hits.append((i, t[:120]))
    if hits:
        for idx, txt in hits:
            flag(label, f'para[{idx}]: {txt}')
    else:
        good(label, 'Not found in body text')

# ─────────────────────────────────────────────────────────────────────────────
# 3. TEST RESULT VALUES
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 3. TEST RESULTS ===')

# Table[9] — summary test results
t9 = doc.tables[9]
checks_t9 = [
    (1, 2, '108',               'T9 Jest total'),
    (1, 3, '108/108 (100%)',    'T9 Jest pass rate'),
    (2, 3, '69/69 (100.0%)',    'T9 Newman pass rate'),
    (3, 2, '213',               'T9 Playwright total'),
    (3, 3, '213/213 (100.0%)', 'T9 Playwright pass rate'),
    (3, 4, '0 skip',            'T9 Playwright skip=0'),
    (3, 4, '0 flaky',           'T9 Playwright flaky=0'),
]
for ri, ci, expected, label in checks_t9:
    actual = t9.rows[ri].cells[ci].text.strip()
    if expected in actual:
        good(label, f'"{expected}" ✓')
    else:
        flag(label, f'expected "{expected}" | actual "{actual[:80]}"')

# Stale numbers in test table
for ri, row in enumerate(t9.rows):
    for ci, cell in enumerate(row.cells):
        ct = cell.text.strip()
        for stale in ['206/211', '97.6%', '211', '206', '3 skip', '2 flaky', '4.8 min']:
            if stale in ct:
                flag(f'T9 stale value', f'R{ri}C{ci} still has "{stale}": {ct[:60]}')

# Table[15] — Playwright per-group
t15 = doc.tables[15]
total_row = None
for ri, row in enumerate(t15.rows):
    if 'TOTAL' in row.cells[0].text or 'รวม' in row.cells[0].text:
        total_row = ri
checks_t15 = [
    (14, 2, '5',   'T15 OAuth Client pass=5'),
    (14, 3, '0',   'T15 OAuth Client skip=0'),
]
if total_row:
    checks_t15 += [
        (total_row, 1, '213', 'T15 TOTAL tests=213'),
        (total_row, 2, '213', 'T15 TOTAL pass=213'),
        (total_row, 3, '0',   'T15 TOTAL skip=0'),
    ]
for ri, ci, expected, label in checks_t15:
    actual = t15.rows[ri].cells[ci].text.strip()
    if expected in actual:
        good(label, f'"{expected}" ✓')
    else:
        flag(label, f'expected "{expected}" | actual "{actual[:60]}"')

# Stale numbers in T15
for ri, row in enumerate(t15.rows):
    for ci, cell in enumerate(row.cells):
        ct = cell.text.strip()
        for stale in ['206', '211', '4.8 min']:
            if stale in ct:
                flag('T15 stale', f'R{ri}C{ci} still has "{stale}": {ct[:60]}')

# ─────────────────────────────────────────────────────────────────────────────
# 4. JWT ALGORITHM
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 4. JWT ALGORITHM ===')
t5 = doc.tables[5]
jwt_checks = [
    (1, 1, 'HS256',         'T5 access token algo'),
    (1, 2, 'HS256',         'T5 refresh token algo'),
    (1, 3, 'Symmetric',     'T5 algo remark'),
    (7, 1, 'JWT_SECRET',    'T5 key source access'),
    (7, 2, 'JWT_SECRET',    'T5 key source refresh'),
    (8, 1, 'N/A',           'T5 public key expose access'),
    (8, 2, 'N/A',           'T5 public key expose refresh'),
]
for ri, ci, expected, label in jwt_checks:
    try:
        actual = t5.rows[ri].cells[ci].text.strip()
        if expected in actual:
            good(label, f'"{expected}" ✓')
        else:
            flag(label, f'expected "{expected}" | actual "{actual[:80]}"')
    except IndexError:
        flag(label, f'Row/col out of range: R{ri}C{ci}')

# ─────────────────────────────────────────────────────────────────────────────
# 5. RATE LIMITING
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 5. RATE LIMITING ===')
t3 = doc.tables[3]
t7 = doc.tables[7]
rl_checks = [
    (t3, 1, 1, '5',      'T3 login max=5'),
    (t3, 2, 1, '5',      'T3 register max=5'),
    (t7, 1, 2, '5',      'T7 login max=5'),
    (t7, 2, 2, '5',      'T7 register max=5'),
    (t7, 9, 1, '1 min',  'T7 userinfo window=1min'),
    (t7, 9, 2, '60',     'T7 userinfo max=60'),
]
for tbl, ri, ci, expected, label in rl_checks:
    try:
        actual = tbl.rows[ri].cells[ci].text.strip()
        if expected in actual:
            good(label, f'"{expected}" ✓')
        else:
            flag(label, f'expected "{expected}" | actual "{actual[:60]}"')
    except IndexError:
        flag(label, f'Row/col out of range')

# ─────────────────────────────────────────────────────────────────────────────
# 6. BCRYPT
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 6. BCRYPT ===')
t8 = doc.tables[8]
actual = t8.rows[6].cells[2].text.strip()
if 'cost-10' in actual:
    good('T8 bcrypt cost-10', f'✓ "{actual[:60]}"')
else:
    flag('T8 bcrypt', f'expected cost-10 | actual "{actual[:60]}"')
if 'cost-12' in actual:
    flag('T8 bcrypt cost-12', f'still has cost-12: "{actual[:60]}"')

# ─────────────────────────────────────────────────────────────────────────────
# 7. TOKEN TYPE LABELS (Table[6])
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 7. TOKEN LABELS ===')
t6 = doc.tables[6]
token_checks = [
    (2, 0, 'Access Token', 'T6 R2 30-day token label'),
    (3, 0, 'Access Token', 'T6 R3 1-hour token label'),
]
for ri, ci, expected, label in token_checks:
    try:
        actual = t6.rows[ri].cells[ci].text.strip()
        if expected in actual:
            good(label, f'"{expected}" ✓')
        else:
            flag(label, f'expected "{expected}" | actual "{actual[:60]}"')
    except IndexError:
        flag(label, 'Row/col out of range')

# ─────────────────────────────────────────────────────────────────────────────
# 8. TABLE NUMBERING (caption cross-refs in body text)
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 8. TABLE NUMBERING CAPTIONS ===')
table_caption_checks = [
    (374, 'ตารางที่ 3.1', 'Requirements table'),
    (641, 'ตารางที่ 3.2', 'API table'),
    (641, '44 endpoint', 'API count 44'),
    (803, 'ตารางที่ 3.3', 'JWT table'),
    (807, 'ตารางที่ 3.4', 'Rate limit table'),
    (821, 'ตารางที่ 3.5', 'PDPA table'),
]
for idx, expected, label in table_caption_checks:
    if idx < len(paras):
        actual = paras[idx].text
        if expected in actual:
            good(f'Caption para[{idx}] {label}', f'"{expected}" ✓')
        else:
            flag(f'Caption para[{idx}] {label}', f'expected "{expected}" | actual "{actual[:100]}"')

# ─────────────────────────────────────────────────────────────────────────────
# 9. API TABLE (Table[2]) — verify fixed routes
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 9. API TABLE ROUTES ===')
t2 = doc.tables[2]
api_checks = [
    # (row, col, expected_text, label)
    (4,  2, '/api/auth/refresh-token',         'R04 refresh-token path'),
    (6,  2, '/api/users/profile',              'R06 PUT profile path'),
    (12, 1, 'POST',                            'R12 change-password method=POST'),
    (12, 2, '/api/auth/change-password',       'R12 change-password path'),
    (13, 2, '/api/auth/validate-token',        'R13 validate-token path'),
    (14, 1, 'DELETE',                          'R14 consents method=DELETE'),
    (14, 2, '/api/oauth/consents/:clientId',   'R14 consents path'),
    (16, 2, '/api/users/export',               'R16 export path'),
    (20, 1, 'POST',                            'R20 session revoke method=POST'),
    (20, 2, '/api/auth/sessions/revoke',       'R20 session revoke path'),
    (21, 1, 'POST',                            'R21 revoke-all method=POST'),
    (21, 2, '/api/auth/sessions/revoke-all-others', 'R21 revoke-all path'),
    (39, 2, '/api/users/:id',                  'R39 users/:id path (no /role)'),
    (44, 2, '/api/oauth/authorize',            'R44 authorize GET path'),
    # Old wrong values should NOT be present
]
# Also check old wrong values are gone
old_wrong = [
    (4,  2, '/api/auth/refresh',    'R04 old refresh path GONE'),
    (12, 1, 'PUT',                  'R12 old PUT GONE'),
    (16, 2, 'export-data',          'R16 old export-data GONE'),
    (20, 2, 'sessions/:id',         'R20 old sessions/:id GONE'),
    (39, 2, '/role',                'R39 old /role GONE'),
    (13, 2, 'resend-verification',  'R13 old resend GONE'),
    (14, 2, 'verify-email',         'R14 old verify-email GONE'),
    (44, 2, 'silent-auth',          'R44 old silent-auth GONE'),
]

for ri, ci, expected, label in api_checks:
    actual = t2.rows[ri].cells[ci].text.strip()
    if expected in actual:
        good(label, f'"{expected}" ✓')
    else:
        flag(label, f'expected "{expected}" | actual "{actual[:80]}"')

for ri, ci, bad_val, label in old_wrong:
    actual = t2.rows[ri].cells[ci].text.strip()
    # For R04: /api/auth/refresh should not appear without -token after it
    if ri == 4 and ci == 2:
        ok = '/api/auth/refresh-token' in actual and '/api/auth/refresh ' not in actual and actual.endswith('refresh-token')
        if ok:
            good(label, 'old short path gone ✓')
        else:
            flag(label, f'old value may still present: "{actual[:60]}"')
    elif ri == 20 and ci == 2:
        # sessions/revoke is correct, sessions/:id is wrong
        ok = '/:id' not in actual
        if ok:
            good(label, '/:id gone ✓')
        else:
            flag(label, f'still has /:id: "{actual[:60]}"')
    else:
        if bad_val not in actual:
            good(label, f'"{bad_val}" gone ✓')
        else:
            flag(label, f'still has "{bad_val}": "{actual[:60]}"')

# ─────────────────────────────────────────────────────────────────────────────
# 10. FIGURE CROSS-REFERENCES
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 10. FIGURE CROSS-REFERENCES ===')
fig_refs = {}
for i, p in enumerate(paras):
    t = p.text
    if not t.strip() or p.style.name.startswith('Heading'): continue
    for ref in re.findall(r'รูปที่\s+(\d+\.\d+)', t):
        fig_refs.setdefault(ref, []).append((i, t[:100]))
if fig_refs:
    print('  Figure refs found in body text:')
    for ref in sorted(fig_refs):
        for idx, txt in fig_refs[ref]:
            print(f'    รูปที่ {ref} at para[{idx}]: {txt[:80]}')
    note('Figure refs', f'{sum(len(v) for v in fig_refs.values())} refs — verify numbering is still correct')
else:
    good('Figure refs', 'None found (or no numbered figures)')

# ─────────────────────────────────────────────────────────────────────────────
# 11. KEY PARAGRAPH CONTENT
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 11. KEY PARAGRAPH CONTENT ===')
key_para_checks = [
    (1007, '213 test cases', 'Playwright 213 in para[1007]'),
    (1007, '0 skip',         'Playwright 0 skip in para[1007]'),
    (1083, '213/213',        'Playwright 213/213 in para[1083]'),
    (1083, '100.0%',         'Playwright 100.0% in para[1083]'),
]
for idx, expected, label in key_para_checks:
    if idx < len(paras):
        actual = paras[idx].text
        if expected in actual:
            good(label, f'✓')
        else:
            flag(label, f'expected "{expected}" | text: "{actual[:120]}"')

# ─────────────────────────────────────────────────────────────────────────────
# 12. HEADING STRUCTURE — count chapters
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 12. HEADING STRUCTURE ===')
h1s = [(i, p.text) for i, p in enumerate(paras) if p.style.name == 'Heading 1']
h2s = [(i, p.text) for i, p in enumerate(paras) if p.style.name == 'Heading 2']
print(f'  Heading 1 count: {len(h1s)}')
for idx, txt in h1s:
    print(f'    [{idx}] {txt[:80]}')
print(f'  Heading 2 count: {len(h2s)}')

# ─────────────────────────────────────────────────────────────────────────────
# 13. PKCE — S256 only
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 13. PKCE ===')
pkce_plain = False
for i, p in enumerate(paras):
    if 'plain' in p.text.lower() and 'pkce' in p.text.lower():
        note(f'PKCE para[{i}]', f'plain mentioned with PKCE: {p.text[:100]}')
        pkce_plain = True
if not pkce_plain:
    good('PKCE plain', 'No "plain" PKCE method in body text')

# Check S256 mentioned correctly
s256_found = any('S256' in p.text for p in paras)
if s256_found:
    good('PKCE S256', 'S256 found in document')
else:
    note('PKCE S256', 'S256 not found — may be in table only')

# ─────────────────────────────────────────────────────────────────────────────
# 14. PDPA SECTION NUMBERS
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 14. PDPA SECTION NUMBERS ===')
pdpa_sections = ['19', '27', '28', '33', '37', '40']
full_text = ' '.join(p.text for p in paras)
for sec in pdpa_sections:
    pattern = rf'มาตรา\s+{sec}\b'
    if re.search(pattern, full_text):
        good(f'PDPA มาตรา {sec}', 'found ✓')
    else:
        note(f'PDPA มาตรา {sec}', 'not found in body — may be table-only')

# ─────────────────────────────────────────────────────────────────────────────
# 15. TABLES COUNT
# ─────────────────────────────────────────────────────────────────────────────
print('\n=== 15. TABLE COUNT ===')
print(f'  Total tables in document: {len(doc.tables)}')
for ti, tbl in enumerate(doc.tables):
    nrows = len(tbl.rows)
    ncols = len(tbl.columns) if tbl.rows else 0
    first_cell = tbl.rows[0].cells[0].text.strip()[:40] if tbl.rows else ''
    print(f'    Table[{ti}]: {nrows} rows x {ncols} cols | first="{first_cell}"')

# ─────────────────────────────────────────────────────────────────────────────
# REPORT
# ─────────────────────────────────────────────────────────────────────────────
print('\n' + '='*70)
print(f'RESULT: {len(issues)} FAIL | {len(ok_list)} OK | {len(notes)} NOTE')
print('='*70)

if issues:
    print('\n--- FAILURES ---')
    for x in issues: print(x)

if notes:
    print('\n--- NOTES (verify manually) ---')
    for x in notes: print(x)

print(f'\n--- OK ({len(ok_list)}) ---')
for x in ok_list: print(x)
