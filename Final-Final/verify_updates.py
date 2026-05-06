import sys
from docx import Document
sys.stdout.reconfigure(encoding='utf-8')

DOC_PATH = r"C:\Users\ASUS\Documents\TEST\TESTMONO\Final-Final\TAS8-final (1).docx"
doc = Document(DOC_PATH)

print("=== Table[9] Summary ===")
for ri, row in enumerate(doc.tables[9].rows):
    cells = [c.text.strip() for c in row.cells]
    print(f"  Row[{ri}]: {cells}")

print("\n=== Table[10] Total row ===")
row28 = doc.tables[10].rows[28]
print(f"  {[c.text.strip() for c in row28.cells]}")

print("\n=== Table[13] ZAP ===")
for ri, row in enumerate(doc.tables[13].rows):
    cells = [c.text.strip() for c in row.cells]
    print(f"  Row[{ri}]: {cells}")

print("\n=== Table[15] TOTAL row ===")
row28 = doc.tables[15].rows[28]
print(f"  {[c.text.strip() for c in row28.cells]}")

print("\n=== Key paragraphs ===")
for i in [979, 980, 990, 1004, 1007, 1026]:
    txt = doc.paragraphs[i].text[:150]
    if txt.strip():
        print(f"  [{i}] {txt}")
